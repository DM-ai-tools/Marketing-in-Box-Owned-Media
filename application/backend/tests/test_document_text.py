"""The clip button's text extraction: real PDF and Word files built in memory, and the refusals."""

from __future__ import annotations

import io
import zipfile

import docx
import pytest
from fastapi.testclient import TestClient

from app.main import app
from app.services import document_text
from app.services.document_text import DocumentTextError, extract


def _pdf(text: str | None) -> bytes:
    """A one-page PDF; `None` gives a page with no text layer (a scan)."""
    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode() if text else b""
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    return out.getvalue()


def _docx() -> bytes:
    document = docx.Document()
    document.add_heading("Skill: Pillar page copy", level=1)
    document.add_paragraph("Lead with the outcome, never the method.")
    document.add_paragraph("One idea per section", style="List Bullet")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Rule"
    table.rows[0].cells[1].text = "Why"
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def test_a_pdf_gives_its_text_and_page_count():
    result = extract("framework.pdf", _pdf("StoryBrand hero framework"))
    assert "StoryBrand hero framework" in result.text
    assert result.kind == "pdf" and result.pages == 1 and not result.truncated


def test_a_word_file_keeps_headings_lists_and_tables():
    result = extract("Skill.DOCX", _docx())
    assert "# Skill: Pillar page copy" in result.text
    assert "Lead with the outcome, never the method." in result.text
    assert "- One idea per section" in result.text
    assert "| Rule | Why |" in result.text
    assert result.kind == "docx"


@pytest.mark.parametrize("name,kind", [("notes.md", "markdown"), ("notes.markdown", "markdown"), ("notes.txt", "text")])
def test_markdown_and_text_are_decoded(name, kind):
    result = extract(name, "# AIDA\r\n\r\n\r\n\r\nAttention first—always.".encode("utf-8"))
    assert result.text == "# AIDA\n\nAttention first—always."
    assert result.kind == kind


def test_cp1252_text_still_reads():
    assert extract("old.txt", b"caf\xe9 \x96 menu").text == "caf\xe9 – menu"


def test_long_text_is_cut_and_says_so():
    result = extract("big.txt", ("line of skill text\n" * 5000).encode())
    assert result.truncated
    assert len(result.text) <= document_text.MAX_CHARS
    assert result.chars > document_text.MAX_CHARS


@pytest.mark.parametrize(
    "name,data,needle",
    [
        ("legacy.doc", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1rest", ".docx or PDF"),
        ("scan.pdf", _pdf(None), "no text layer"),
        ("fake.pdf", b"hello", "not a PDF"),
        ("fake.docx", b"hello", "not a Word document"),
        ("renamed.docx", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1rest", "old Word (.doc)"),
        ("image.png", b"\x89PNG", "PDF, Word (.docx), Markdown or text"),
        ("empty.txt", b"", "empty"),
        ("blank.txt", b"   \n\n ", "No text"),
        ("hidden.txt", b"%PDF-1.4 binary", "binary document"),
    ],
)
def test_unreadable_files_are_refused_with_the_fix(name, data, needle):
    with pytest.raises(DocumentTextError, match=needle.replace("(", r"\(").replace(")", r"\)").replace(".", r"\.")):
        extract(name, data)


def test_oversized_files_are_refused(monkeypatch):
    monkeypatch.setattr(document_text, "MAX_UPLOAD_BYTES", 10)
    with pytest.raises(DocumentTextError, match="over"):
        extract("a.txt", b"x" * 11)


def test_a_zip_bomb_docx_is_refused(monkeypatch):
    monkeypatch.setattr(document_text, "MAX_UNPACKED_BYTES", 1000)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "0" * 5000)
    with pytest.raises(DocumentTextError, match="too large once unpacked"):
        extract("bomb.docx", out.getvalue())


def test_the_route_returns_text_and_422s_with_the_sentence():
    client = TestClient(app)
    ok = client.post("/pipeline/files/extract", files={"file": ("skill.md", b"# Skill\n\nBe brief.", "text/markdown")})
    assert ok.status_code == 200
    assert ok.json()["text"] == "# Skill\n\nBe brief."
    assert ok.json()["filename"] == "skill.md"
    bad = client.post("/pipeline/files/extract", files={"file": ("old.doc", b"\xd0\xcf\x11\xe0", "application/msword")})
    assert bad.status_code == 422
    assert ".docx or PDF" in bad.json()["detail"]
